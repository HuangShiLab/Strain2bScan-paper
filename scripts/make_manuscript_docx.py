#!/usr/bin/env python3
"""Build the editable DOCX snapshot from the assembled manuscript.

Pandoc performs Markdown/DOCX conversion. python-docx is optional; when
available it tightens table layout and writes document metadata. Install it
with `python3 -m pip install python-docx` for fully styled output.
"""
from __future__ import annotations
from pathlib import Path
import re
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript/full_manuscript.md"
OUT = ROOT / "manuscript/Strain2bScan_manuscript_figures_tables_2026-10-13.docx"
TITLE = "Strain2bScan: strain-level 2bRAD-marker profiling of microbiomes"


def page_break_openxml() -> str:
    return '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'


def image_size(path: Path) -> tuple[int, int]:
    out = subprocess.check_output(["sips", "-g", "pixelWidth", "-g", "pixelHeight", str(path)], text=True)
    return (int(re.search(r"pixelWidth:\s*(\d+)", out).group(1)),
            int(re.search(r"pixelHeight:\s*(\d+)", out).group(1)))


def prepare_intermediate() -> Path:
    text = SRC.read_text()
    legends_start = text.index("## Figure legends")
    tables_start = text.index("## Tables", legends_start)
    references_start = text.index("## References", tables_start)
    body = text[:legends_start].rstrip()
    legends = text[legends_start:tables_start]
    tables = text[tables_start:references_start].rstrip()
    references = text[references_start:].rstrip()

    captions = {m.group(2): " ".join(m.group(1).split()) for m in re.finditer(
        r"(^\*\*Figure (S?\d+)\..*?)(?=^\*\*Figure |\Z)", legends, flags=re.M | re.S)}

    def image_block(label: str) -> str:
        pattern = f"Fig{int(label):02d}_*.png" if label.isdigit() else f"Fig{label}_*.png"
        matches = list((ROOT / "figures/numbered").glob(pattern))
        if not matches:
            raise FileNotFoundError(f"missing numbered figure for {label}: {pattern}")
        path = matches[0]
        width_px, height_px = image_size(path)
        width_in = min(6.3, 7.8 * width_px / height_px)
        return "\n\n".join([
            f"``` {{=openxml}}\n{page_break_openxml()}\n```",
            f"![]({path.relative_to(ROOT)}){{width={width_in:.2f}in}}",
            captions[label],
        ])

    parts = [body, "# Main figures", ""]
    parts.extend(image_block(str(i)) for i in range(1, 12))
    parts.extend(["# Supplementary figures", ""])
    parts.extend(image_block(f"S{i}") for i in range(1, 6))
    table_lines = []
    for line in tables.splitlines():
        if line.startswith("**Table "):
            table_lines.extend(["``` {=openxml}", page_break_openxml(), "```", ""])
        table_lines.append(line)
    parts.extend(["\n".join(table_lines).strip(), references])
    intermediate = Path("/tmp/Strain2bScan_final_docx.md")
    intermediate.write_text("\n\n".join(parts) + "\n")
    return intermediate


def style_docx(path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_LINE_SPACING
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError:
        print("python-docx unavailable; retained Pandoc table styling")
        return

    doc = Document(path)
    for section in doc.sections:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    def set_child(parent, tag: str, **attrs):
        element = parent.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            parent.append(element)
        for key, value in attrs.items():
            element.set(qn(key), str(value))
        return element

    def style_table(table):
        tbl_pr = table._tbl.tblPr
        set_child(tbl_pr, "w:tblW", **{"w:w": "5000", "w:type": "pct"})
        set_child(tbl_pr, "w:tblLayout", **{"w:type": "autofit"})
        old_borders = tbl_pr.find(qn("w:tblBorders"))
        if old_borders is not None:
            tbl_pr.remove(old_borders)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge_el = OxmlElement(f"w:{edge}")
            edge_el.set(qn("w:val"), "single")
            edge_el.set(qn("w:sz"), "4")
            edge_el.set(qn("w:space"), "0")
            edge_el.set(qn("w:color"), "B7C4D6")
            borders.append(edge_el)
        tbl_pr.append(borders)
        old_margins = tbl_pr.find(qn("w:tblCellMar"))
        if old_margins is not None:
            tbl_pr.remove(old_margins)
        margins = OxmlElement("w:tblCellMar")
        for side, width in (("top", "20"), ("left", "55"), ("bottom", "20"), ("right", "55")):
            element = OxmlElement(f"w:{side}")
            element.set(qn("w:w"), width)
            element.set(qn("w:type"), "dxa")
            margins.append(element)
        tbl_pr.append(margins)
        for tc_pr in tbl_pr.iter(qn("w:tcPr")):
            tcw = tc_pr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tc_pr.append(tcw)
            tcw.set(qn("w:w"), "0")
            tcw.set(qn("w:type"), "auto")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        ncols = len(table.columns)
        body_size = Pt(6.8) if ncols >= 10 else Pt(7.6) if ncols >= 7 else Pt(8.4)
        header_size = Pt(7.2) if ncols >= 10 else Pt(8.0) if ncols >= 7 else Pt(8.8)
        for row_number, row in enumerate(table.rows, start=1):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            if row_number == 1 and tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                if row_number == 1:
                    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                    if shd is None:
                        shd = OxmlElement("w:shd")
                        cell._tc.get_or_add_tcPr().append(shd)
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "DCE6F1")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    for run in paragraph.runs:
                        run.font.size = header_size if row_number == 1 else body_size
                        run.font.bold = row_number == 1

    for table in doc.tables:
        style_table(table)

    metadata = doc.core_properties
    metadata.title = TITLE
    metadata.author = "Yufeng Zhang; Shi Huang"
    metadata.subject = "Manuscript snapshot with embedded main and supplementary figures"
    metadata.keywords = "Strain2bScan; 2bRAD; metagenomics; strain profiling"
    metadata.comments = "Nature-style optimized manuscript snapshot 2026-10-13"
    doc.save(path)


def reorder_ooxml(path: Path) -> None:
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError:
        return
    orders = {
        "w:tblPr": ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                    "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
                    "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"],
        "w:trPr": ["w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter", "w:cantSplit",
                   "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange"],
        "w:tcPr": ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap",
                   "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers", "w:cellIns",
                   "w:cellDel", "w:cellMerge", "w:tcPrChange"],
    }

    def reorder(parent, order_tag: str):
        rank = {qn(tag): i for i, tag in enumerate(orders[order_tag])}
        for child in sorted(parent, key=lambda e: rank.get(e.tag, len(rank))):
            parent.append(child)

    doc = Document(path)
    for table in doc.tables:
        reorder(table._tbl.tblPr, "w:tblPr")
        for row in table.rows:
            if row._tr.trPr is not None:
                reorder(row._tr.trPr, "w:trPr")
            for cell in row.cells:
                if cell._tc.tcPr is not None:
                    reorder(cell._tc.tcPr, "w:tcPr")
    doc.save(path)


def main() -> None:
    intermediate = prepare_intermediate()
    subprocess.run([
        "pandoc", str(intermediate), "--from", "markdown+raw_attribute+pipe_tables",
        "--to", "docx", "--output", str(OUT), "--resource-path",
        f"{ROOT}:{ROOT/'figures'}",
    ], check=True)
    style_docx(OUT)
    reorder_ooxml(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
